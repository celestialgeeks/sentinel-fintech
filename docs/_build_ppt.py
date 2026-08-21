"""
Slide 1: Title Slide
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RgbColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Slide 1: Title
blank = prs.slide_layouts[6]  # blank
slide = prs.slides.add_slide(blank)

# Dark navy background
background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
background.fill.solid()
background.fill.fore_color.rgb = RgbColor(15, 52, 96)
background.line.fill.background()

# Title text
title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(12.333), Inches(1.5))
tf = title_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "SENTINEL"
p.font.size = Pt(54)
p.font.bold = True
p.font.color.rgb = RgbColor(255, 255, 255)
p.alignment = PP_ALIGN.CENTER

# Subtitle
sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(12.333), Inches(1.0))
tf = sub_box.text_frame
p = tf.paragraphs[0]
p.text = "AI-Powered Real-Time Transaction Screening & Fraud Detection"
p.font.size = Pt(24)
p.font.color.rgb = RgbColor(100, 200, 255)
p.alignment = PP_ALIGN.CENTER

# Bottom info
info_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(12.333), Inches(1.5))
tf = info_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "B.Tech Major Project-I | CSE-AIML Department | IPS Academy, Indore (RGPV University)"
p.font.size = Pt(14)
p.font.color.rgb = RgbColor(180, 180, 200)
p.alignment = PP_ALIGN.CENTER
p = tf.add_paragraph()
p.text = "Academic Year 2026-27"
p.font.size = Pt(14)
p.font.color.rgb = RgbColor(180, 180, 200)
p.alignment = PP_ALIGN.CENTER

print("Slide 1 done")
prs.save("/tmp/sentinel_ppt_temp.pptx")
print("Saved temp")
